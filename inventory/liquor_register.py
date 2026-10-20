"""
The liquor stock register: each day, each bottle, from Rasova's stock ledger.

A bar keeps a daily stock register for the excise department: what it had,
what came in (from KSBCL, with the invoice), what it sold, what broke, and
what it has left, by brand and bottle size. Rasova already records every one
of those movements (inventory.InventoryTransaction): a delivery is "restock",
a peg sold is "consume" (a cancel puts it back), breakage is "wastage", and a
bar count's correction is "adjustment". This lays them out per business day.

The official layout of the Karnataka register wasn't found published; this
is the plain form (opening + received - sold - broken +/- counted = closing,
in bottles, ml and bulk litres) to be matched to a pub's accountant's
columns.

Run: python manage.py test inventory.tests.test_liquor_register
"""
from bisect import bisect_right
from datetime import timedelta
from decimal import Decimal


from core.utils import get_business_date_range
from inventory.bar_count import bar_items
from inventory.models import InventoryTransaction
from inventory.unit_conversion import convert_quantity

MAX_DAYS = 62

KINDS = {"restock": "received", "consume": "sold", "wastage": "broken", "adjustment": "counted"}


def _ml(quantity, item):
    """A quantity of `item` in ml, or None if it can't be said (stock kept
    in pieces with no bottle size)."""
    if item.unit in ("ml", "l"):
        return convert_quantity(Decimal(quantity), item.unit, "ml")
    if item.unit == "pcs" and item.bottle_ml:
        return Decimal(quantity) * item.bottle_ml
    return None


def register_rows(tenant, outlet, start_date, end_date):
    """One row per business day per bar item, oldest first. Two queries
    however many items and days: the items, and every one of their ledger
    rows from the first day on, each put in its day by a binary search. The
    stock at the start = today's stock minus everything recorded since."""
    days = []
    day = start_date
    while day <= end_date and len(days) < MAX_DAYS:
        days.append((day, *get_business_date_range(day, outlet)))
        day += timedelta(days=1)
    if not days:
        return []
    starts = [start for _, start, _ in days]
    last_end = days[-1][2]

    items = list(bar_items(tenant, outlet))
    since = {item.id: Decimal("0") for item in items}          # everything since the first day began
    moves = {item.id: [{kind: Decimal("0") for kind in KINDS.values()} for _ in days] for item in items}
    refs = {item.id: [[] for _ in days] for item in items}
    ledger = (InventoryTransaction.objects.filter(item_id__in=list(since), created_at__gte=starts[0])
              .values_list("item_id", "created_at", "quantity", "transaction_type", "reference"))
    for item_id, created_at, quantity, kind, reference in ledger:
        since[item_id] += quantity
        if created_at >= last_end:
            continue
        index = bisect_right(starts, created_at) - 1
        moves[item_id][index][KINDS.get(kind, "counted")] += quantity
        if kind == "restock" and reference:
            refs[item_id][index].append(reference)

    rows = []
    for item in items:
        stock = item.stock - since[item.id]
        for index, (day, _, _) in enumerate(days):
            moved = moves[item.id][index]
            opening = stock
            stock = opening + sum(moved.values(), Decimal("0"))
            row = {
                "day": day, "item": item.name, "unit": item.unit, "bottle_ml": item.bottle_ml,
                "opening": opening, "received": moved["received"], "sold": -moved["sold"],
                "broken": -moved["broken"], "counted": moved["counted"], "closing": stock,
                "references": ", ".join(refs[item.id][index]),
            }
            for key in ("opening", "received", "sold", "broken", "counted", "closing"):
                ml = _ml(row[key], item)
                row[f"{key}_ml"] = ml
                row[f"{key}_bl"] = (ml / 1000).quantize(Decimal("0.001")) if ml is not None else None
            if item.bottle_ml and row["closing_ml"] is not None:
                row["closing_bottles"] = int(row["closing_ml"] // item.bottle_ml)
                row["closing_loose_ml"] = row["closing_ml"] - row["closing_bottles"] * item.bottle_ml
            rows.append(row)
    rows.sort(key=lambda r: (r["day"], r["item"]))
    return rows


HEADERS = [
    ("Day", "day"), ("Brand / item", "item"), ("Bottle (ml)", "bottle_ml"),
    ("Opening (bulk litres)", "opening_bl"), ("Received (bulk litres)", "received_bl"),
    ("Sold (bulk litres)", "sold_bl"), ("Broken (bulk litres)", "broken_bl"),
    ("Count correction (bulk litres)", "counted_bl"), ("Closing (bulk litres)", "closing_bl"),
    ("Closing: full bottles", "closing_bottles"), ("Closing: loose ml", "closing_loose_ml"),
    ("Received: invoice / reference", "references"),
]


def register_workbook(tenant, outlet, start_date, end_date):
    """The register as an Excel workbook (openpyxl)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Liquor stock register"
    ws.append([f"{outlet.name}: liquor stock register, {start_date:%d %b %Y} to {end_date:%d %b %Y}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Opening + received - sold - broken +/- count correction = closing. "
               "Match these columns to your excise register before filing."])
    ws.append([label for label, _ in HEADERS])
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for row in register_rows(tenant, outlet, start_date, end_date):
        ws.append([row.get(key) if key != "day" else row["day"].strftime("%d-%m-%Y") for _, key in HEADERS])
    for col, (label, _) in zip("ABCDEFGHIJKL", HEADERS):
        ws.column_dimensions[col].width = max(12, min(40, len(label) + 2))
    ws.freeze_panes = "A4"
    return wb
